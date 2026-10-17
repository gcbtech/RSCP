"""
Expected-package imports: CSV/XLSX uploads, email ingest, and the one-time
legacy manifest.csv import.

The `packages` table is the only source of truth. Every import writes
straight into it (inside the caller's transaction) and leaves a row in
`package_imports` so the dashboard can show when data last came in.

Identity of a package row: (tracking_number, item_name). One shipment
(tracking number) can carry several different items, each its own row;
receiving a tracking number receives all of them (see log_receipt).

Placeholders: when the real carrier tracking is not known yet a row uses an
order ID instead, either 'ORDER-<order id>[-NN]' (created by email ingest) or
the raw order ID (an upload row with no tracking). When the real tracking
arrives, from an upload or a shipping email, the placeholder row is updated
in place, so its id, priority flag, manual date and SKU survive.
"""
import csv
import datetime
import difflib
import io
import json
import logging
import re

from app.utils.helpers import parse_date

logger = logging.getLogger(__name__)

PLACEHOLDER_PREFIX = 'ORDER-'
SOURCE_UPLOAD = 'manifest'   # packages.source of uploaded rows (historic value, shown as MANIFEST)
SOURCE_EMAIL = 'email'
SOURCE_MANUAL = 'manual'

# Statuses that are still "on the way" and derive from the expected date.
OPEN_STATUSES = ('on_time', 'expected', 'past_due', 'pending', 'incoming')
# Statuses an import must never overwrite or merge into.
CLOSED_STATUSES = ('return_pending', 'returned', 'refunded')

AMAZON_ORDER_RE = re.compile(r'^\d{3}-\d{7}-\d{7}$')
_ORDER_IN_PLACEHOLDER_RE = re.compile(r'^(\d{3}-\d{7}-\d{7}|\d{2}-\d{5}-\d{5})(?:-\d{2})?$')
_EBAY_DUP_TRACKING_RE = re.compile(r'^trackingnumber\.\d+$', re.IGNORECASE)

ORDER_ID_HEADERS = ['order id', 'order number', 'order #', 'reference number',
                    'reference #', 'ref number', 'order']
NAME_MATCH_THRESHOLD = 0.2   # email and upload titles for the same item can differ wildly


class ImportFileError(ValueError):
    """The uploaded file cannot be read as a table."""


# ---------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------

def _dedupe_headers(raw_headers):
    """Strip headers and make them unique the way pandas did ('A', 'A.1', ...),
    so column names (and saved habits like picking 'TrackingNumber.1') carry over."""
    out, seen = [], set()
    for i, h in enumerate(raw_headers):
        name = str(h).strip() if h is not None else ''
        if not name:
            name = f'Unnamed: {i}'
        candidate, n = name, 0
        while candidate in seen:
            n += 1
            candidate = f'{name}.{n}'
        seen.add(candidate)
        out.append(candidate)
    return out


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))   # Excel stores 1234567890 and qty 2 as floats
    return str(value).strip()


def _finish_table(raw_rows):
    """First row is the header. Drops blank rows and trailing empty columns."""
    if not raw_rows:
        return [], []
    width = max(len(r) for r in raw_rows)
    rows = [list(r) + [''] * (width - len(r)) for r in raw_rows]
    last = -1
    for i in range(width):
        if any(r[i] != '' for r in rows):
            last = i
    rows = [r[:last + 1] for r in rows]
    headers = _dedupe_headers(rows[0])
    body = [r for r in rows[1:] if any(c != '' for c in r)]
    return headers, body


def read_csv_bytes(data):
    if data[:2] == b'PK':
        raise ImportFileError("This looks like an Excel file. Save it as .xlsx or CSV and try again.")
    if b'\0' in data[:1024]:
        raise ImportFileError("This file is binary, not CSV. Save it as standard CSV and try again.")
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('cp1252', errors='replace')
    raw = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text, newline=''))]
    return _finish_table(raw)


def read_xlsx_bytes(data):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Could not open the Excel file: {e}")
    try:
        ws = wb.worksheets[0]   # same sheet pandas.read_excel used
        ws.reset_dimensions()   # some exporters write a wrong used-range; read every cell
        raw = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return _finish_table(raw)


def read_table(data, filename):
    """Return (headers, rows) where rows are lists of strings."""
    if (filename or '').lower().endswith('.xlsx'):
        return read_xlsx_bytes(data)
    return read_csv_bytes(data)


def save_table(path, headers, rows):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump({'headers': headers, 'rows': rows}, f)


def load_table(path):
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    return data['headers'], data['rows']


# ---------------------------------------------------------------------------
# Rows -> normalized items
# ---------------------------------------------------------------------------

def _clean_id(value):
    """Tracking / order IDs: drop Excel ="..." wrapping and quotes."""
    s = str(value or '').replace('="', '').replace('"', '').strip()
    return '' if s.lower() in ('nan', 'none', 'null') else s


def _clean_name(value):
    # Same normalization the old manifest sync used, so re-importing a file
    # matches the rows it created.
    s = str(value or '').strip().replace('="', '').replace('"', '')
    return '' if s.lower() in ('nan', 'none') else s


def _clean_text(value):
    s = str(value or '').strip()
    return '' if s.lower() in ('nan', 'none') else s


def _parse_qty(value):
    try:
        return int(float(value or 1))
    except (TypeError, ValueError):
        return 1


def normalize_date(value):
    s = str(value or '').strip()
    m = re.match(r'^(\d{4}-\d{2}-\d{2})[ T]', s)
    if m:
        s = m.group(1)
    return parse_date(s)


def status_for_date(date_str, today=None):
    today = today or datetime.date.today()
    try:
        d = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return 'on_time'
    if d == today:
        return 'expected'
    if d < today:
        return 'past_due'
    return 'on_time'


def order_id_from_placeholder(tracking):
    """'ORDER-113-1234567-1234567-01' -> '113-1234567-1234567' (else None)."""
    t = tracking or ''
    if not t.upper().startswith(PLACEHOLDER_PREFIX):
        return None
    m = _ORDER_IN_PLACEHOLDER_RE.match(t[len(PLACEHOLDER_PREFIX):])
    return m.group(1) if m else None


def items_from_mapping(headers, rows, mapping):
    """Turn table rows into package items using a column mapping
    {'tracking','name','date','quantity','asin','image','product_url'} -> header.

    Returns (items, skipped) where skipped counts rows without any tracking.
    """
    idx = {h: i for i, h in enumerate(headers)}

    def col(row, key):
        h = mapping.get(key)
        if not h or h not in idx:
            return ''
        i = idx[h]
        return row[i] if i < len(row) else ''

    lower = {h.lower(): h for h in headers}
    track_col = mapping.get('tracking')
    order_cols = [lower[h] for h in ORDER_ID_HEADERS if h in lower and lower[h] != track_col]

    # eBay exports have two 'TrackingNumber' columns: the first holds the
    # order ID, the duplicate ('TrackingNumber.1') the carrier tracking.
    ebay_base = lower.get('trackingnumber')
    ebay_dups = [h for h in headers if _EBAY_DUP_TRACKING_RE.match(h)]
    ebay = bool(ebay_base and ebay_dups and track_col in [ebay_base] + ebay_dups)

    items, skipped = [], 0
    for row in rows:
        def cell(h):
            i = idx[h]
            return row[i] if i < len(row) else ''

        order_id = ''
        if ebay:
            order_id = _clean_id(cell(ebay_base))
            tracking = ''
            for h in ebay_dups:
                tracking = _clean_id(cell(h))
                if tracking:
                    break
            if len(order_id) <= 5:
                order_id = ''
        else:
            tracking = _clean_id(col(row, 'tracking'))
            for h in order_cols:
                v = _clean_id(cell(h))
                if len(v) > 4:
                    order_id = v
                    break

        placeholder = False
        if not tracking and order_id:
            tracking, placeholder = order_id, True
        elif tracking and AMAZON_ORDER_RE.match(tracking):
            # An Amazon order number in the tracking column: not shippable yet
            order_id, placeholder = order_id or tracking, True
        elif tracking and tracking == order_id:
            placeholder = True
        if not tracking:
            skipped += 1
            continue

        asin = _clean_name(col(row, 'asin')).strip()
        image = _clean_text(col(row, 'image'))
        if not image and len(asin) >= 10:
            image = f"https://images-na.ssl-images-amazon.com/images/P/{asin.upper()}.01._SX200_.jpg"

        items.append({
            'tracking': tracking,
            'order_id': order_id or None,
            'placeholder': placeholder,
            'item_name': _clean_name(col(row, 'name')) or 'Unknown',
            'date': normalize_date(col(row, 'date')),
            'quantity': _parse_qty(col(row, 'quantity')),
            'image_url': image,
            'asin': asin,
            'source_url': _clean_text(col(row, 'product_url')),
        })
    return _dedupe_items(items), skipped


def _dedupe_items(items):
    """One item per (tracking, name), last row wins (as the old sync did).
    A placeholder row is dropped when the same file already has that item
    with real tracking."""
    by_key = {}
    for it in items:
        by_key[(it['tracking'], it['item_name'])] = it
    shipped = {(it['order_id'], it['item_name']) for it in by_key.values()
               if it['order_id'] and not it['placeholder']}
    return [it for it in by_key.values()
            if not (it['placeholder'] and (it['order_id'], it['item_name']) in shipped)]


# Header aliases the old manifest sync understood (exact names).
_LEGACY_ALIASES = {
    'tracking': ['TrackingNumber', 'Carrier Tracking #', 'Tracking Number', 'Tracking'],
    'name': ['ItemName', 'Title', 'Product Name', 'Item Description'],
    'date': ['Date', 'Order Date', 'Purchase Date', 'Time'],
    'quantity': ['Quantity', 'Item Quantity', 'Qty', 'Order Quantity'],
    'image': ['Image', 'ImageUrl', 'Photo', 'Image URL'],
    'asin': ['ASIN', 'Item ID', 'ItemID'],
    'product_url': ['SourceURL', 'ProductURL', 'URL', 'PurchaseURL', 'Link', 'ProductLink',
                    'Product URL', 'View Order Detail'],
}


def legacy_mapping(headers):
    present = set(headers)
    return {key: next((a for a in aliases if a in present), None)
            for key, aliases in _LEGACY_ALIASES.items()}


# ---------------------------------------------------------------------------
# Writing to the database (callers own the transaction)
# ---------------------------------------------------------------------------

def _name_ratio(a, b):
    return difflib.SequenceMatcher(None, (a or '').lower(), (b or '').lower()).ratio()


def _best_match(rows, name, claimed=(), fallback_first=False):
    rows = [r for r in rows if r['id'] not in claimed]
    if not rows:
        return None
    if len(rows) == 1:
        return rows[0]
    best = max(rows, key=lambda r: _name_ratio(name, r['item_name']))
    if _name_ratio(name, best['item_name']) >= NAME_MATCH_THRESHOLD:
        return best
    return rows[0] if fallback_first else None


def _like_escape(s):
    return s.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')


def is_placeholder_row(row):
    t = row['tracking_number'] or ''
    return t.upper().startswith(PLACEHOLDER_PREFIX) or (bool(row['order_id']) and t == row['order_id'])


def _order_rows(conn, order_id, unscanned_only):
    sql = """
        SELECT * FROM packages
         WHERE (order_id = ? OR tracking_number = ? OR tracking_number = ?
                OR tracking_number LIKE ? ESCAPE '\\')
           AND COALESCE(status, '') NOT IN ('return_pending', 'returned', 'refunded')
    """
    if unscanned_only:
        sql += " AND date_scanned IS NULL AND COALESCE(source, '') != 'manual'"
    return conn.execute(sql + " ORDER BY id", (
        order_id, order_id, PLACEHOLDER_PREFIX + order_id,
        PLACEHOLDER_PREFIX + _like_escape(order_id) + '-%')).fetchall()


def find_placeholder(conn, order_id, item_name, tracking=None):
    """Best unscanned row of this order that an imported item should merge
    into: a placeholder, or an email row already carrying the same tracking
    (email and file titles for one item rarely match). None if there is none."""
    rows = [r for r in _order_rows(conn, order_id, unscanned_only=True)
            if is_placeholder_row(r)
            or (tracking and r['tracking_number'] == tracking and r['source'] == SOURCE_EMAIL)]
    return _best_match(rows, item_name)


def _mapped_sku(conn, item_name):
    row = conn.execute("SELECT inventory_sku FROM product_mappings WHERE package_name = ?",
                       (item_name,)).fetchone()
    return row['inventory_sku'] if row else None


def _apply_changes(conn, row, changes):
    """UPDATE only the columns whose value actually changes. Returns True if any did."""
    changes = {k: v for k, v in changes.items() if row[k] != v}
    if not changes:
        return False
    sets = ', '.join(f'{k} = ?' for k in changes)
    conn.execute(f'UPDATE packages SET {sets} WHERE id = ?', list(changes.values()) + [row['id']])
    return True


def _update_existing(conn, row, item, today):
    """Re-import of a known (tracking, name). Received and manually added
    packages are locked: only blank fields get filled."""
    locked = row['date_scanned'] is not None or row['source'] == SOURCE_MANUAL \
        or row['status'] in CLOSED_STATUSES
    changes = {}
    for col, key in (('image_url', 'image_url'), ('asin', 'asin'), ('source_url', 'source_url'),
                     ('order_id', 'order_id')):
        val = item.get(key)
        if val and (not row[col] or not locked):
            changes[col] = val
    if not locked:
        date_final = row['manual_date'] or item['date']
        if date_final == 'Pending' and row['date_expected']:
            date_final = row['date_expected']   # a file without dates must not wipe them
        changes['date_expected'] = date_final
        changes['quantity'] = item['quantity']
        if row['status'] in OPEN_STATUSES:
            changes['status'] = status_for_date(date_final, today)
    if not row['sku']:
        sku = _mapped_sku(conn, row['item_name'])
        if sku:
            changes['sku'] = sku
    return _apply_changes(conn, row, changes)


def _merge_into_placeholder(conn, ph, item, source, today):
    """Give a placeholder row the data of an imported item. A real tracking
    number replaces the placeholder's; an order-ID-only item just enriches it."""
    changes = {}
    if not item['placeholder']:
        changes['tracking_number'] = item['tracking']
        changes['item_name'] = item['item_name']
        changes['source'] = source
    if item.get('order_id'):
        changes['order_id'] = item['order_id']
    if not ph['manual_date'] and item['date'] != 'Pending':
        changes['date_expected'] = item['date']
    date_final = changes.get('date_expected', ph['date_expected'])
    if ph['status'] in OPEN_STATUSES:
        changes['status'] = status_for_date(date_final, today)
    # Trust the email's quantity when it found one; otherwise the file's.
    if not (ph['source'] == SOURCE_EMAIL and (ph['quantity'] or 0) > 1):
        changes['quantity'] = item['quantity']
    for col in ('image_url', 'asin', 'source_url'):
        if item.get(col):
            changes[col] = item[col]
    if not ph['sku']:
        sku = _mapped_sku(conn, changes.get('item_name', ph['item_name']))
        if sku:
            changes['sku'] = sku
    return _apply_changes(conn, ph, changes)


def _insert(conn, item, source, today):
    conn.execute('''
        INSERT INTO packages (tracking_number, item_name, date_expected, quantity, image_url,
                              status, asin, source_url, sku, source, order_id)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (item['tracking'], item['item_name'], item['date'], item['quantity'],
          item.get('image_url') or '', status_for_date(item['date'], today),
          item.get('asin') or '', item.get('source_url') or '',
          _mapped_sku(conn, item['item_name']), source, item.get('order_id')))


def _get_row(conn, tracking, item_name):
    return conn.execute("SELECT * FROM packages WHERE tracking_number = ? AND item_name = ?",
                        (tracking, item_name)).fetchone()


def apply_items(conn, items, source=SOURCE_UPLOAD, insert_only=False, today=None):
    """Upsert normalized items into packages. Returns (added, updated).

    insert_only: leave rows that already exist untouched (legacy import).
    """
    today = today or datetime.date.today()
    added = updated = 0
    for item in items:
        existing = _get_row(conn, item['tracking'], item['item_name'])
        if existing:
            if not insert_only and _update_existing(conn, existing, item, today):
                updated += 1
            continue
        if item.get('order_id'):
            ph = find_placeholder(conn, item['order_id'], item['item_name'],
                                  None if item['placeholder'] else item['tracking'])
            if ph:
                if _merge_into_placeholder(conn, ph, item, source, today):
                    updated += 1
                continue
            if item['placeholder']:
                # The order already shipped under a real tracking number
                known = [r for r in _order_rows(conn, item['order_id'], unscanned_only=False)
                         if not is_placeholder_row(r) and r['item_name'] == item['item_name']]
                if known:
                    continue
        _insert(conn, item, source, today)
        added += 1
    return added, updated


def ingest_email_items(conn, items, today=None):
    """Write parsed email orders (email_ingest.check_amazon_emails) into
    packages. Returns (added, updated).

    An email about an order we already know (from an upload, an earlier email,
    or a received package) enriches that row instead of adding a duplicate; a
    shipping email that carries the carrier tracking gives it to the order's
    placeholder row.
    """
    today = today or datetime.date.today()
    added = updated = 0
    claimed = set()
    for raw in items:
        placeholder_tracking = _clean_id(raw.get('tracking'))
        order_id = _clean_id(raw.get('order_id')) or order_id_from_placeholder(placeholder_tracking)
        carrier = _clean_id(raw.get('carrier_tracking')).upper()
        name = _clean_name(raw.get('name')) or 'Unknown'
        qty = _parse_qty(raw.get('quantity'))
        image = _clean_text(raw.get('image_url'))
        date = normalize_date(raw.get('date'))
        tracking = carrier or placeholder_tracking
        if not tracking:
            continue

        row = _get_row(conn, tracking, name)
        if row:
            claimed.add(row['id'])
            changes = {}
            if image and not row['image_url']:
                changes['image_url'] = image
            if order_id and not row['order_id']:
                changes['order_id'] = order_id
            if _apply_changes(conn, row, changes):
                updated += 1
            continue

        if order_id:
            match = _best_match(_order_rows(conn, order_id, unscanned_only=False), name,
                                claimed, fallback_first=True)
            if match:
                claimed.add(match['id'])
                changes = {}
                if not match['order_id']:
                    changes['order_id'] = order_id
                if image and not match['image_url']:
                    changes['image_url'] = image
                open_row = match['date_scanned'] is None and match['source'] != SOURCE_MANUAL
                if open_row and qty > 1:
                    changes['quantity'] = qty
                if open_row and carrier and is_placeholder_row(match) \
                        and not _get_row(conn, carrier, match['item_name']):
                    changes['tracking_number'] = carrier
                if _apply_changes(conn, match, changes):
                    updated += 1
                continue

        conn.execute('''
            INSERT INTO packages (tracking_number, item_name, date_expected, quantity, status,
                                  image_url, source, order_id, sku)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (tracking, name, date, qty, status_for_date(date, today), image, SOURCE_EMAIL,
              order_id or None, _mapped_sku(conn, name)))
        row = _get_row(conn, tracking, name)
        if row:
            claimed.add(row['id'])
        added += 1
    return added, updated


def record_import(conn, source, filename, added, updated, user):
    conn.execute('''
        INSERT INTO package_imports (source, filename, rows_added, rows_updated, user)
        VALUES (?, ?, ?, ?, ?)
    ''', (source, filename, int(added), int(updated), user))


def last_import(conn):
    """Most recent import with its age in hours, or None."""
    row = conn.execute('''
        SELECT source, filename, rows_added, rows_updated, imported_at, user,
               (julianday('now') - julianday(imported_at)) * 24.0 AS age_hours
          FROM package_imports ORDER BY imported_at DESC, id DESC LIMIT 1
    ''').fetchone()
    return dict(row) if row else None


def refresh_package_statuses(conn):
    """Recompute on_time / expected / past_due from the expected date (store's
    local calendar) for packages not yet received. Email rows carry the order
    date, not a delivery estimate, so they are left alone. Returns rows changed."""
    new_status = """CASE WHEN date_expected = date('now', 'localtime') THEN 'expected'
                         WHEN date_expected < date('now', 'localtime') THEN 'past_due'
                         ELSE 'on_time' END"""
    cur = conn.execute(f'''
        UPDATE packages SET status = {new_status}
         WHERE date_scanned IS NULL
           AND status IN ('on_time', 'expected', 'past_due', 'pending')
           AND COALESCE(source, '') != 'email'
           AND date_expected GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]'
           AND status != {new_status}
    ''')
    return cur.rowcount


def trim_old_packages(conn):
    """AUTO_TRIM: drop received, imported packages expected over 60 days ago."""
    cur = conn.execute('''
        DELETE FROM packages
         WHERE date_scanned IS NOT NULL
           AND source IN ('manifest', 'email')
           AND date_expected GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]'
           AND date_expected < date('now', 'localtime', '-60 days')
    ''')
    return cur.rowcount


# ---------------------------------------------------------------------------
# Entry points that own their transaction
# ---------------------------------------------------------------------------

def _connect():
    from app.services.db import get_db_connection
    return get_db_connection()


def import_items(items, source, filename, user, package_source=SOURCE_UPLOAD):
    """Upsert items and log the import, all in one write transaction.
    Returns (added, updated)."""
    conn = _connect()
    try:
        conn.execute('BEGIN IMMEDIATE')
        try:
            added, updated = apply_items(conn, items, source=package_source)
            record_import(conn, source, filename, added, updated, user)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        return added, updated
    finally:
        conn.close()


def _current_username():
    try:
        from flask import has_request_context
        from flask_login import current_user
        if has_request_context() and current_user.is_authenticated:
            return current_user.username
    except Exception:
        pass
    return None


def run_email_ingest(fetch=None):
    """The one email-ingest path (background job and the admin button).

    Returns {'status': 'success'|'skipped'|'error', 'message', 'count',
    'added', 'updated'}; 'failed': True when something raised.
    """
    from app.services.data_manager import load_config
    try:
        config = load_config()
        if not config.get('EMAIL_INGEST_ENABLED', False):
            return {'status': 'skipped', 'message': 'Email Ingest Disabled', 'count': 0}
        server = config.get('IMAP_SERVER')
        user = config.get('EMAIL_USER')
        password = config.get('EMAIL_PASS') or config.get('EMAIL_PASSWORD')
        if not (server and user and password):
            return {'status': 'error', 'message': 'Missing Credentials', 'count': 0}

        if fetch is None:
            import email_ingest
            fetch = email_ingest.check_amazon_emails
        logger.info(f"[Email Ingest] Checking {user} on {server}...")
        items = fetch(server, user, password)
        if not items:
            return {'status': 'success', 'message': 'No new emails found', 'count': 0,
                    'added': 0, 'updated': 0}

        conn = _connect()
        try:
            conn.execute('BEGIN IMMEDIATE')
            try:
                added, updated = ingest_email_items(conn, items)
                if added or updated:
                    record_import(conn, 'email', None, added, updated,
                                  _current_username() or 'system')
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        finally:
            conn.close()
        logger.info(f"[Email Ingest] {added} added, {updated} updated.")
        return {'status': 'success', 'count': added,
                'message': f"Imported {added} new, updated {updated}",
                'added': added, 'updated': updated}
    except Exception as e:
        logger.error(f"[Email Ingest] Failed: {e}")
        return {'status': 'error', 'message': str(e), 'count': 0, 'failed': True}


def run_package_maintenance():
    """Periodic job: keep statuses in step with the calendar; AUTO_TRIM."""
    from app.services.data_manager import load_config
    conn = _connect()
    try:
        conn.execute('BEGIN IMMEDIATE')
        try:
            changed = refresh_package_statuses(conn)
            trimmed = trim_old_packages(conn) if load_config().get('AUTO_TRIM', False) else 0
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    finally:
        conn.close()
    return changed, trimmed


# ---------------------------------------------------------------------------
# One-time legacy manifest.csv import (schema migration 10)
# ---------------------------------------------------------------------------

def import_legacy_manifest(conn, path):
    """Insert rows of an old manifest.csv that are not in packages yet.
    Idempotent; the file itself is left alone. Returns rows added."""
    with open(path, 'rb') as f:
        data = f.read()
    headers, rows = read_csv_bytes(data)
    if not headers:
        return 0
    items, _ = items_from_mapping(headers, rows, legacy_mapping(headers))
    # Only bring in recent/upcoming packages. Older CSV rows that are missing
    # from the database were usually removed on purpose (AUTO_TRIM deleted
    # rows but left them in the CSV) and would come back as stale "past due".
    import datetime as _dt
    cutoff = (_dt.date.today() - _dt.timedelta(days=30)).strftime('%Y-%m-%d')
    recent = [i for i in items if i['date'] == 'Pending' or i['date'] >= cutoff]
    if len(recent) < len(items):
        logger.info(f"Legacy manifest import: skipped {len(items) - len(recent)} rows dated before {cutoff}")
    added, _ = apply_items(conn, recent, source=SOURCE_UPLOAD, insert_only=True)
    if added:
        record_import(conn, 'upload', 'manifest.csv (legacy import)', added, 0, 'system')
    return added
